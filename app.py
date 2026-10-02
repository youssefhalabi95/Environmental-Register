from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    send_file
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO


app = Flask(__name__)

# Used by Flask to securely sign the user's session
app.secret_key = "environmental-register-secret-key"

# Free version limit
MAX_ASPECTS = 5


# --------------------------------------------------
# RISK CALCULATION
# --------------------------------------------------

def get_risk_level(score):

    if score <= 4:
        return "Low"

    elif score <= 9:
        return "Medium"

    elif score <= 16:
        return "High"

    else:
        return "Critical"


# --------------------------------------------------
# LANDING PAGE
# --------------------------------------------------

@app.route("/")
def landing():

    return render_template("landing.html")


# --------------------------------------------------
# ENVIRONMENTAL REGISTER
# --------------------------------------------------

@app.route("/register", methods=["GET", "POST"])
def register():

    project = session.get("project", {})
    register_data = session.get("register", [])

    message = None

    if request.method == "POST":

        form_type = request.form.get("form_type")

        # ------------------------------------------
        # SAVE PROJECT INFORMATION
        # ------------------------------------------

        if form_type == "project":

            project = {
                "name": request.form["project_name"],
                "client": request.form["client"],
                "location": request.form["location"],
                "type": request.form["project_type"],
                "prepared_by": request.form["prepared_by"],
                "date": request.form["assessment_date"]
            }

            session["project"] = project

            return redirect(url_for("register"))


        # ------------------------------------------
        # ADD NEW ENVIRONMENTAL ASPECT
        # ------------------------------------------

        elif form_type == "aspect":

            if len(register_data) >= MAX_ASPECTS:

                message = (
                    f"The free version allows a maximum of "
                    f"{MAX_ASPECTS} environmental aspects."
                )

            else:

                activity = request.form["activity"]
                aspect = request.form["aspect"]
                impact = request.form["impact"]
                condition = request.form["condition"]

                severity = int(request.form["severity"])
                likelihood = int(request.form["likelihood"])

                score = severity * likelihood
                level = get_risk_level(score)

                entry = {
                    "activity": activity,
                    "aspect": aspect,
                    "impact": impact,
                    "condition": condition,
                    "severity": severity,
                    "likelihood": likelihood,
                    "score": score,
                    "level": level
                }

                register_data.append(entry)

                session["register"] = register_data

                return redirect(url_for("register"))


    return render_template(
        "index.html",
        project=project,
        register=register_data,
        message=message,
        max_aspects=MAX_ASPECTS
    )


# --------------------------------------------------
# EDIT ASPECT
# --------------------------------------------------

@app.route("/edit/<int:index>", methods=["GET", "POST"])
def edit(index):

    register_data = session.get("register", [])

    # Make sure the requested row exists
    if index < 0 or index >= len(register_data):

        return redirect(url_for("register"))


    if request.method == "POST":

        activity = request.form["activity"]
        aspect = request.form["aspect"]
        impact = request.form["impact"]
        condition = request.form["condition"]

        severity = int(request.form["severity"])
        likelihood = int(request.form["likelihood"])

        score = severity * likelihood
        level = get_risk_level(score)

        register_data[index] = {
            "activity": activity,
            "aspect": aspect,
            "impact": impact,
            "condition": condition,
            "severity": severity,
            "likelihood": likelihood,
            "score": score,
            "level": level
        }

        session["register"] = register_data

        return redirect(url_for("register"))


    return render_template(
        "edit.html",
        item=register_data[index],
        index=index
    )


# --------------------------------------------------
# DELETE ASPECT
# --------------------------------------------------

@app.route("/delete/<int:index>")
def delete(index):

    register_data = session.get("register", [])

    if 0 <= index < len(register_data):

        register_data.pop(index)

        session["register"] = register_data

    return redirect(url_for("register"))


# --------------------------------------------------
# CLEAR EVERYTHING
# --------------------------------------------------

@app.route("/clear")
def clear():

    session.pop("project", None)
    session.pop("register", None)

    return redirect(url_for("register"))


# --------------------------------------------------
# EXCEL EXPORT
# --------------------------------------------------

@app.route("/export")
def export_excel():

    project = session.get("project", {})
    register_data = session.get("register", [])


    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Environmental Register"


    # ------------------------------------------
    # TITLE
    # ------------------------------------------

    worksheet["A1"] = "ENVIRONMENTAL ASPECT & IMPACT REGISTER"

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet.merge_cells("A1:I1")


    # ------------------------------------------
    # PROJECT INFORMATION
    # ------------------------------------------

    worksheet["A3"] = "Project Information"

    worksheet["A3"].font = Font(
        bold=True,
        size=12
    )


    project_information = [
        ("Project Name", project.get("name", "")),
        ("Client", project.get("client", "")),
        ("Project Location", project.get("location", "")),
        ("Project Type", project.get("type", "")),
        ("Prepared By", project.get("prepared_by", "")),
        ("Assessment Date", project.get("date", ""))
    ]


    row = 4

    for label, value in project_information:

        worksheet.cell(row=row, column=1).value = label
        worksheet.cell(row=row, column=2).value = value

        worksheet.cell(row=row, column=1).font = Font(
            bold=True
        )

        row += 1


    # ------------------------------------------
    # REGISTER TITLE
    # ------------------------------------------

    worksheet["A11"] = "Environmental Aspect & Impact Register"

    worksheet["A11"].font = Font(
        bold=True,
        size=12
    )


    # ------------------------------------------
    # HEADERS
    # ------------------------------------------

    headers = [
        "No.",
        "Activity",
        "Environmental Aspect",
        "Environmental Impact",
        "Condition",
        "Severity",
        "Likelihood",
        "Risk Score",
        "Risk Level"
    ]


    for column, header in enumerate(headers, start=1):

        cell = worksheet.cell(
            row=12,
            column=column
        )

        cell.value = header

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="1F4D3A"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    # ------------------------------------------
    # REGISTER DATA
    # ------------------------------------------

    for number, item in enumerate(register_data, start=1):

        row = number + 12

        values = [
            number,
            item["activity"],
            item["aspect"],
            item["impact"],
            item["condition"],
            item["severity"],
            item["likelihood"],
            item["score"],
            item["level"]
        ]


        for column, value in enumerate(values, start=1):

            worksheet.cell(
                row=row,
                column=column
            ).value = value


    # ------------------------------------------
    # COLUMN WIDTHS
    # ------------------------------------------

    widths = {
        "A": 8,
        "B": 25,
        "C": 30,
        "D": 30,
        "E": 15,
        "F": 12,
        "G": 12,
        "H": 12,
        "I": 15
    }


    for column, width in widths.items():

        worksheet.column_dimensions[column].width = width


    # Freeze header row
    worksheet.freeze_panes = "A13"


    # Add Excel filters
    if register_data:

        worksheet.auto_filter.ref = (
            f"A12:I{12 + len(register_data)}"
        )


    # Wrap text
    for row_cells in worksheet.iter_rows():

        for cell in row_cells:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


    # ------------------------------------------
    # CREATE DOWNLOAD
    # ------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)


    return send_file(
        output,
        as_attachment=True,
        download_name="environmental_register.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# --------------------------------------------------
# ABOUT
# --------------------------------------------------

@app.route("/about")
def about():

    return render_template("about.html")


# --------------------------------------------------
# PRIVACY POLICY
# --------------------------------------------------

@app.route("/privacy")
def privacy():

    return render_template("privacy.html")


# --------------------------------------------------
# RUN APPLICATION
# --------------------------------------------------

if __name__ == "__main__":

    app.run(debug=True)